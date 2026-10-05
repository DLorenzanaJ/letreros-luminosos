"""Genera el reporte Word a partir de salidas/metrics.json (los números siempre coinciden)."""
import json
import os
import sys
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from calidad.requisitos import REQUISITOS

SAL, FIG = "salidas", "salidas/figuras"
AZUL, VERDE, NARANJA, GRIS, ROJO = "#1F4E79", "#2E8B57", "#E07B00", "#6B7280", "#C0392B"
CONFIG_DEFECTO = {"institucion": "[Nombre de tu universidad]", "materia": "[Nombre de la materia]",
                  "alumno": "[Tu nombre completo]", "matricula": "[Tu matrícula]",
                  "profesor": "[Nombre del profesor]", "grupo": "[Grupo]",
                  "fecha": "Octubre de 2026", "repositorio": "https://github.com/[tu-usuario]/letreros-luminosos"}


# ----------------------------------------------------------------- figuras
def _guardar(nombre):
    ruta = f"{FIG}/{nombre}.png"
    plt.tight_layout()
    plt.savefig(ruta, dpi=150)
    plt.close()
    return ruta


def figura_terminal(archivo, nombre, titulo, ancho=104, max_lineas=70):
    with open(f"{SAL}/{archivo}", encoding="utf-8") as f:
        lineas = []
        for linea in f.read().splitlines():
            lineas += textwrap.wrap(linea, ancho, subsequent_indent="    ") or [""]
    lineas = [l.replace("$", r"\$") for l in lineas[:max_lineas]]
    alto = len(lineas) * 0.155 + 0.7
    plt.figure(figsize=(9.2, alto), facecolor="#1e1e1e")
    plt.axis("off")
    plt.text(0.0, 1.0, f"PS C:\\letreros> {titulo}", color="#7ec699", family="monospace", fontsize=7.5, va="top",
             transform=plt.gca().transAxes)
    plt.text(0.0, 1.0 - 0.45 / alto, "\n".join(lineas), color="#d4d4d4", family="monospace", fontsize=7,
             va="top", transform=plt.gca().transAxes, linespacing=1.3)
    plt.savefig(f"{FIG}/{nombre}.png", dpi=150, facecolor="#1e1e1e", bbox_inches="tight")
    plt.close()
    return f"{FIG}/{nombre}.png"


def _caja(ax, x, y, w, h, texto, color):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02", fc=color, ec="white", lw=1.5))
    ax.text(x + w / 2, y + h / 2, texto, ha="center", va="center", color="white", fontsize=8.5, weight="bold")


def _flecha(ax, x1, y1, x2, y2):
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", color="#333", lw=1.6))


def figura_pipeline():
    fig, ax = plt.subplots(figsize=(9.5, 3.2))
    ax.axis("off"); ax.set_xlim(0, 10); ax.set_ylim(0, 3.2)
    pasos = [("Desarrollador\n(VS Code + IA)", AZUL), ("git push\nGitHub", GRIS), ("GitHub Actions\n(CI)", NARANJA),
             ("pytest +\ncobertura", VERDE), ("Métricas +\nrevisión IA", VERDE), ("Reporte\n.docx", AZUL)]
    for i, (texto, color) in enumerate(pasos):
        _caja(ax, 0.1 + i * 1.65, 1.7, 1.35, 0.9, texto, color)
        if i < len(pasos) - 1:
            _flecha(ax, 1.5 + i * 1.65, 2.15, 1.75 + i * 1.65, 2.15)
    _caja(ax, 3.4, 0.2, 1.7, 0.8, "Docker\n(ejecución reproducible)", GRIS)
    _caja(ax, 5.4, 0.2, 1.7, 0.8, "Artefactos\n(metrics.json)", GRIS)
    _flecha(ax, 4.0, 1.7, 4.2, 1.02)
    _flecha(ax, 6.3, 1.7, 6.2, 1.02)
    return _guardar("pipeline")


def figura_arquitectura():
    fig, ax = plt.subplots(figsize=(9.5, 4.2))
    ax.axis("off"); ax.set_xlim(0, 10); ax.set_ylim(0, 4.2)
    _caja(ax, 3.9, 3.3, 2.2, 0.7, "main.py (consola)", AZUL)
    for i, nombre in enumerate(["catalogo", "cotizador", "ventas", "soporte_ia"]):
        _caja(ax, 0.4 + i * 2.4, 1.9, 1.9, 0.7, nombre, VERDE)
        _flecha(ax, 5.0, 3.3, 1.35 + i * 2.4, 2.62)
    _caja(ax, 1.2, 0.2, 2.8, 0.7, "repositorio JSON\n(Fase 2: SQLite)", GRIS)
    _caja(ax, 6.0, 0.2, 3.2, 0.7, "Proveedor IA: Claude (API)\no modo local", NARANJA)
    _flecha(ax, 5.6, 1.9, 2.6, 0.92)
    _flecha(ax, 8.0, 1.9, 7.8, 0.92)
    return _guardar("arquitectura")


def figura_complejidad(m):
    f = list(reversed(m["producto"]["top_funciones"]))
    plt.figure(figsize=(8, 3.6))
    plt.barh([f"{x['modulo']}.{x['funcion']}" for x in f], [x["cc"] for x in f], color=AZUL)
    plt.axvline(10, color=ROJO, ls="--", label="Límite recomendado (10)")
    plt.xlabel("Complejidad ciclomática"); plt.legend()
    return _guardar("complejidad")


def figura_cobertura(m):
    c = m["cobertura"]["por_archivo"]
    plt.figure(figsize=(8, 3.4))
    barras = plt.bar(list(c), [v["cobertura"] for v in c.values()], color=VERDE)
    plt.axhline(80, color=ROJO, ls="--", label="Meta 80 %")
    plt.bar_label(barras, fmt="%.1f%%"); plt.ylim(0, 110); plt.ylabel("Cobertura (%)"); plt.legend(loc="lower right")
    return _guardar("cobertura")


def figura_defectos(m):
    d = m["defectos"]
    fig, ejes = plt.subplots(1, 2, figsize=(9, 3.4))
    ejes[0].bar(list(d["por_fase"]), list(d["por_fase"].values()), color=[AZUL, VERDE, ROJO])
    ejes[0].set_title("Defectos por fase de detección")
    mods = d["por_modulo"]
    ejes[1].bar(list(mods), [v["densidad"] for v in mods.values()], color=NARANJA)
    ejes[1].set_title("Densidad por módulo (def/KLOC)"); ejes[1].tick_params(axis="x", rotation=30)
    return _guardar("defectos")


def figura_tiempos(m):
    b = m["defectos"]["bitacora"]
    plt.figure(figsize=(9, 3.4))
    x = range(len(b))
    plt.bar([i - 0.2 for i in x], [r["mttd"] for r in b], 0.4, label="Tiempo de detección (h)", color=AZUL)
    plt.bar([i + 0.2 for i in x], [r["mttr"] for r in b], 0.4, label="Tiempo de reparación (h)", color=NARANJA)
    plt.xticks(list(x), [r["id"] for r in b]); plt.ylabel("Horas"); plt.legend()
    return _guardar("tiempos")


def figura_desviacion(m):
    mods = m["proyecto"]["modulos"]
    plt.figure(figsize=(9, 3.5))
    x = range(len(mods))
    plt.bar([i - 0.2 for i in x], [r["estimadas"] for r in mods], 0.4, label="Estimadas", color=GRIS)
    plt.bar([i + 0.2 for i in x], [r["reales"] for r in mods], 0.4, label="Reales", color=AZUL)
    plt.xticks(list(x), [r["modulo"] for r in mods], rotation=20); plt.ylabel("Horas"); plt.legend()
    return _guardar("desviacion")


def figura_estimaciones(m):
    c = m["estimaciones"]["comparativo"]
    plt.figure(figsize=(8, 3.5))
    barras = plt.bar([x["tecnica"] for x in c], [x["horas"] for x in c], color=[AZUL, NARANJA, VERDE, GRIS])
    plt.axhline(c[0]["real"], color=ROJO, ls="--", label=f"Horas reales ({c[0]['real']:.0f} h)")
    plt.bar_label(barras, fmt="%.1f"); plt.ylabel("Horas"); plt.legend(); plt.xticks(rotation=10)
    return _guardar("estimaciones")


def captura_o_marcador(nombre, texto):
    for ext in ("png", "jpg", "jpeg"):
        ruta = f"capturas/{nombre}.{ext}"
        if os.path.exists(ruta):
            return ruta
    plt.figure(figsize=(9, 3))
    plt.axis("off")
    plt.gca().add_patch(plt.Rectangle((0, 0), 1, 1, fill=False, ec=GRIS, ls="--", lw=2, transform=plt.gca().transAxes))
    plt.text(0.5, 0.5, f"Pega aquí tu captura real\n(guárdala como capturas/{nombre}.png y vuelve a ejecutar)\n\n{texto}",
             ha="center", va="center", fontsize=10, color=GRIS, transform=plt.gca().transAxes)
    return _guardar(f"marcador_{nombre}")


# ----------------------------------------------------------------- documento
class Reporte:
    def __init__(self):
        self.doc = Document()
        self.nf = self.nt = 0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, lado, Cm(2.5))
        est = self.doc.styles
        est["Normal"].font.name, est["Normal"].font.size = "Calibri", Pt(11)
        for nivel, tam in ((1, 18), (2, 14), (3, 12)):
            e = est[f"Heading {nivel}"]
            e.font.name, e.font.size, e.font.bold = "Calibri", Pt(tam), True
            e.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        est["Heading 1"].paragraph_format.page_break_before = True

    # texto
    def h(self, texto, nivel=1):
        self.doc.add_heading(texto, nivel)

    def p(self, texto, negrita=False, cursiva=False, centrado=False, tam=None, espacio=6):
        par = self.doc.add_paragraph()
        r = par.add_run(texto)
        r.bold, r.italic = negrita, cursiva
        if tam:
            r.font.size = Pt(tam)
        par.paragraph_format.space_after = Pt(espacio)
        if centrado:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return par

    def lista(self, items):
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def formula(self, texto):
        par = self.doc.add_paragraph()
        r = par.add_run(texto)
        r.italic = True
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # tabla
    def tabla(self, titulo, enc, filas, anchos, alin=None):
        self.nt += 1
        cap = self.doc.add_paragraph(f"Tabla {self.nt}. {titulo}", style="Caption")
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=len(enc))
        t.style, t.alignment, t.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
        alin = alin or ["l"] * len(enc)
        for i, texto in enumerate(enc):
            self._celda(t.rows[0].cells[i], texto, anchos[i], True, "c", "1F4E79")
        for n, fila in enumerate(filas):
            celdas = t.add_row().cells
            for i, valor in enumerate(fila):
                self._celda(celdas[i], str(valor), anchos[i], False, alin[i], "EAF1F8" if n % 2 else None)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _celda(self, celda, texto, ancho, cabecera, alin, fondo):
        celda.width = Cm(ancho)
        par = celda.paragraphs[0]
        par.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}[alin]
        run = par.add_run(texto)
        run.font.size = Pt(9)
        run.bold = cabecera
        if cabecera:
            run.font.color.rgb = RGBColor(255, 255, 255)
        if fondo:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fondo)
            celda._tc.get_or_add_tcPr().append(shd)

    # figura
    def figura(self, ruta, titulo, explicacion, ancho=15.5):
        self.nf += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(ruta, width=Cm(ancho))
        cap = self.doc.add_paragraph(f"Figura {self.nf}. {titulo}", style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.add_run("Explicación: ").bold = True
        par.add_run(explicacion)

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _campo(self, par, instruccion, texto=""):
        for tipo in ("begin", None, "separate", "text", "end"):
            r = par.add_run()
            if tipo is None:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instruccion
            elif tipo == "text":
                r.text = texto
                continue
            else:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tipo)
            r._r.append(el)

    def portada(self, cfg, m):
        for _ in range(2):
            self.p("", espacio=18)
        self.p(cfg["institucion"].upper(), True, centrado=True, tam=16)
        self.p(cfg["materia"], False, centrado=True, tam=13, espacio=40)
        self.p("Sistema de Ventas de Letreros Luminosos con Soporte de Inteligencia Artificial", True,
               centrado=True, tam=24, espacio=14)
        self.p("Reporte de calidad de software en un flujo DevOps", False, True, True, 14, 50)
        for etiqueta, clave in (("Alumno", "alumno"), ("Matrícula", "matricula"), ("Grupo", "grupo"),
                                ("Profesor", "profesor"), ("Repositorio", "repositorio"), ("Fecha", "fecha")):
            self.p(f"{etiqueta}: {cfg[clave]}", centrado=True, tam=12, espacio=4)
        self.p("", espacio=30)
        self.p(f"Métricas generadas el {m['generado'].replace('T', ' ')} en {m['sistema']} "
               f"(Python {m['python']}). Asistente IA en modo: {m['modo_ia']}.", cursiva=True, centrado=True, tam=9)

    def indice(self):
        self.salto()
        self.p("Índice", True, tam=20)
        par = self.doc.add_paragraph()
        self._campo(par, 'TOC \\o "1-2" \\h \\z \\u',
                    "Al abrir en Word, acepte actualizar los campos (o clic derecho > Actualizar campo) para ver el índice con páginas.")
        for ajuste in (self.doc.settings.element,):
            el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true"); ajuste.append(el)
        pie = self.doc.sections[0].footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pie.add_run("Página ")
        self._campo(pie, "PAGE", "1")


def si_no(x):
    return "Sí" if x else "No"


def generar(m):
    os.makedirs(FIG, exist_ok=True)
    if not os.path.exists("reporte_config.json"):
        with open("reporte_config.json", "w", encoding="utf-8") as f:
            json.dump(CONFIG_DEFECTO, f, ensure_ascii=False, indent=2)
    with open("reporte_config.json", encoding="utf-8") as f:
        cfg = {**CONFIG_DEFECTO, **json.load(f)}
    pr, d, py, co, rq = m["producto"], m["defectos"], m["proyecto"], m["cobertura"], m["requisitos"]
    es, ca, ri = m["estimaciones"], m["calidad"], m["revision_ia"]
    r = Reporte()
    r.portada(cfg, m)
    r.indice()

    # 1 ------------------------------------------------------------------
    r.h("1. Introducción")
    r.p("Este reporte documenta el desarrollo y la validación de un sistema de ventas de letreros luminosos "
        "(neón LED, letras de acrílico con luz y cajas de luz) que incorpora un asistente de soporte con "
        "inteligencia artificial. El proyecto se construyó con prácticas DevOps: control de versiones en GitHub, "
        "integración continua con GitHub Actions, pruebas automáticas, ejecución reproducible con Docker y "
        "automatización del cálculo de métricas y de este mismo reporte.")
    r.p("Una característica central del trabajo es que la documentación coincide con la ejecución del programa: "
        "todas las cifras de este documento (pruebas, cobertura, complejidad, defectos, tiempos y estimaciones) "
        "se leen del archivo salidas/metrics.json, que el programa genera cada vez que se ejecuta el comando "
        f"python run_all.py. La última ejecución corrió {pr['funciones_total']} funciones medidas en "
        f"{pr['sloc']} líneas de código fuente y {m['pruebas']['total']} pruebas automáticas.")
    r.h("1.1 Objetivo general", 2)
    r.p("Desarrollar un sistema de ventas con soporte de IA y validar su calidad mediante un flujo DevOps "
        "automatizado, midiendo producto, proceso y proyecto.")
    r.h("1.2 Objetivos específicos", 2)
    r.lista(["Implementar los módulos de catálogo, cotización, ventas y soporte con IA.",
             "Automatizar pruebas, cobertura, métricas y reporte con un solo comando y con GitHub Actions.",
             "Calcular métricas de producto, proceso y proyecto aplicando IA en la revisión e interpretación.",
             "Aplicar cuatro técnicas de estimación y compararlas contra el esfuerzo real.",
             "Emitir un informe de calidad basado en ISO/IEC 25010 (idoneidad funcional y fiabilidad)."])
    r.h("1.3 Contenido del documento", 2)
    r.p("El documento describe el proyecto y su justificación DevOps, los requisitos del usuario y del negocio, "
        "el aseguramiento de la calidad, las características de calidad evaluadas, la ejecución del programa con "
        "su salida, las métricas, las estimaciones, el informe de calidad y las conclusiones. Cada imagen incluye "
        "una explicación de lo que muestra.")

    # 2 ------------------------------------------------------------------
    r.h("2. Descripción del proyecto y justificación DevOps")
    r.h("2.1 El negocio", 2)
    r.p("La empresa vende letreros luminosos a cafeterías, comercios y talleres. El precio depende del tipo de "
        "letrero y del área en metros cuadrados; además se ofrecen colores RGB (recargo de 15 %) e instalación. "
        "Los pedidos grandes reciben descuento por volumen (5 % desde $10,000 y 10 % desde $25,000, antes de IVA) "
        "y se calcula IVA de 16 %. Los clientes también reportan fallas postventa, que el asistente de IA atiende.")
    r.h("2.2 Justificación: por qué es un proyecto DevOps", 2)
    r.tabla("Prácticas DevOps aplicadas", ["Práctica", "Cómo se aplica en el proyecto", "Evidencia"], [
        ["Control de versiones", "Repositorio Git en GitHub con commits por módulo", "Repositorio (sección 14)"],
        ["Integración continua", "GitHub Actions ejecuta pruebas y métricas en cada push", ".github/workflows/ci.yml"],
        ["Automatización", "Un solo comando ejecuta demo, pruebas, métricas y reporte", "run_all.py"],
        ["Validación de software", "pytest con cobertura y trazabilidad requisito-prueba", "Secciones 4 y 8"],
        ["Contenedores", "Dockerfile para ejecutar igual en cualquier equipo", "Dockerfile"],
        ["Medición continua", "Métricas de producto, proceso y proyecto en JSON", "salidas/metrics.json"],
        ["IA aplicada", "Asistente de soporte y revisión de código con IA", "app/soporte_ia.py, calidad/revision_ia.py"]],
        [3.6, 8.0, 4.9])
    r.h("2.3 Arquitectura", 2)
    r.figura(figura_arquitectura(), "Arquitectura del sistema",
             "main.py es la interfaz de consola y consume cuatro módulos. Ningún módulo toca archivos directamente: "
             "todos usan el repositorio, que hoy guarda JSON y en la Fase 2 se cambiará por SQLite sin modificar la "
             "lógica. El módulo soporte_ia usa un proveedor de IA intercambiable: Claude o Gemini si hay API key, o el modo "
             "local de reglas si no la hay o si la API falla.")
    filas = [[k, v["sloc"], v["funciones"], v["cc_promedio"], v["cc_maxima"]] for k, v in pr["modulos"].items()]
    filas.append(["TOTAL", pr["sloc"], pr["funciones_total"], pr["cc_promedio"], pr["cc_maxima"]])
    r.tabla("Módulos del sistema (medidos con radon)", ["Módulo", "Líneas (SLOC)", "Funciones", "CC promedio", "CC máxima"],
            filas, [4.5, 3.2, 3.0, 3.0, 2.8], ["l", "c", "c", "c", "c"])
    r.h("2.4 Flujo DevOps", 2)
    r.figura(figura_pipeline(), "Flujo de integración continua",
             "El desarrollador sube cambios a GitHub; GitHub Actions ejecuta automáticamente las pruebas con "
             "cobertura, calcula las métricas, corre la revisión con IA y genera el reporte. Docker permite repetir la "
             "misma ejecución en cualquier equipo y los artefactos (metrics.json y reporte) quedan guardados.")
    r.h("2.5 Tecnologías", 2)
    r.tabla("Tecnologías utilizadas", ["Tecnología", "Uso"], [
        ["Python 3.12", "Lenguaje del sistema"], ["pytest / pytest-cov", "Pruebas automáticas y cobertura"],
        ["radon", "Complejidad ciclomática, líneas de código e índice de mantenibilidad"],
        ["python-docx / matplotlib", "Generación del reporte y de las gráficas"],
        ["GitHub / GitHub Actions", "Repositorio e integración continua"], ["Docker", "Ejecución reproducible"],
        ["Claude (Anthropic) o Gemini (Google)", "IA del asistente de soporte, la revisión y la interpretación (opcional)"]],
        [5.0, 11.5])

    # 3 ------------------------------------------------------------------
    r.h("3. Requisitos del usuario y expectativas del negocio")
    r.h("3.1 Requisitos del usuario", 2)
    r.tabla("Necesidades de los usuarios", ["ID", "Usuario", "Necesidad", "Prioridad"], [
        ["U-01", "Cliente", "Conocer los tipos de letrero y compararlos", "Alta"],
        ["U-02", "Cliente", "Obtener una cotización rápida escribiendo las medidas", "Alta"],
        ["U-03", "Cliente", "Reportar una falla y recibir pasos de solución inmediatos", "Alta"],
        ["U-04", "Vendedor", "Registrar pedidos con descuento y seguir su estado", "Alta"],
        ["U-05", "Técnico de soporte", "Ver tickets abiertos por severidad y cerrarlos con tiempo registrado", "Media"],
        ["U-06", "Administrador", "Conocer la calidad del sistema con métricas confiables", "Media"]],
        [1.6, 3.2, 9.5, 2.2], ["c", "l", "l", "c"])
    r.h("3.2 Expectativas del negocio", 2)
    r.tabla("Expectativas del negocio e indicadores", ["Expectativa", "Indicador", "Meta"], [
        ["Cotizar sin errores de cálculo", "Pruebas de cotización aprobadas", "100 %"],
        ["Atender postventa con rapidez", "MTTR de defectos", "<= 24 h"],
        ["Evitar fallas en producción", "Eficacia de pruebas sobre defectos críticos", ">= 70 %"],
        ["Código mantenible", "Complejidad ciclomática máxima", "<= 10"],
        ["Entregar a tiempo", "Desviación de esfuerzo total", "<= 30 %"],
        ["Confiar en los cambios", "Cobertura de código", ">= 80 %"]], [5.5, 7.5, 3.5], ["l", "l", "c"])

    # 4 ------------------------------------------------------------------
    r.h("4. Requisitos explícitos y grado de cumplimiento")
    r.p("Cada requisito está ligado a una o más pruebas automáticas (calidad/requisitos.py). Un requisito se "
        "considera cumplido cuando todas sus pruebas aprobaron en la ejecución actual.")
    r.tabla("Trazabilidad requisito - prueba", ["ID", "Requisito", "Tipo", "Pruebas", "¿Cumple?"],
            [[x["id"], x["requisito"], x["tipo"], x["pruebas"], si_no(x["cumple"])] for x in rq["detalle"]],
            [1.6, 8.6, 2.2, 1.9, 2.2], ["c", "l", "c", "c", "c"])
    r.p(f"Resultado: {rq['cumplidos']} de {rq['total']} requisitos cumplidos, es decir {rq['porcentaje']} %.", True)
    r.tabla("Requisitos del trabajo académico y dónde se cubren", ["Requisito del trabajo", "Sección"], [
        ["Portada, introducción, índice, imágenes explicadas, conclusión", "Todo el documento"],
        ["Programa DevOps con IA, GitHub, ejecución y salida", "2, 7 y 14"],
        ["Calidad de software y grado de cumplimiento; informe de calidad", "4 y 13"],
        ["Requisitos del usuario y expectativas del negocio", "3"],
        ["Aseguramiento de la calidad; enfoque preventivo; carácter sistemático", "5"],
        ["Idoneidad funcional y fiabilidad", "6"],
        ["Métricas de producto, proceso y proyecto", "8, 9 y 10"],
        ["Estimación: expertos, análoga, tres puntos y puntos de función", "11"]], [11.5, 5.0], ["l", "c"])

    # 5 ------------------------------------------------------------------
    r.h("5. Aseguramiento de la calidad de software")
    r.p("El aseguramiento de la calidad (SQA) es el conjunto de actividades planificadas que dan confianza en que "
        "el software cumplirá los requisitos. No se limita a probar al final: se integra en todo el ciclo de desarrollo.")
    r.h("5.1 Enfoque preventivo", 2)
    r.p("El enfoque preventivo busca evitar defectos antes de que existan, en lugar de encontrarlos tarde. "
        "En este proyecto se aplica así:")
    r.tabla("Actividades preventivas", ["Actividad", "Cuándo ocurre", "Herramienta"], [
        ["Requisitos con prueba asociada", "Antes de programar", "calidad/requisitos.py"],
        ["Validación de entradas (medidas, cliente, estados)", "En el diseño de cada función", "Excepciones ValueError"],
        ["Revisión de código con IA y reglas", "En cada ejecución, antes de liberar", "calidad/revision_ia.py"],
        ["Pruebas automáticas en cada push", "Integración continua", "GitHub Actions + pytest"],
        ["Umbrales de calidad (cobertura, complejidad)", "Al calcular métricas", "Informe de calidad"],
        ["Tolerancia a fallos del proveedor de IA", "En el diseño del módulo", "Modo local de respaldo"]],
        [6.2, 5.0, 5.3])
    r.p(f"El efecto de la prevención se observa en que {d['por_fase']['revision']} de {d['total']} defectos "
        f"({d['eficacia_revision']} %) se detectaron en revisión, antes de ejecutar pruebas.")
    r.h("5.2 Carácter sistemático", 2)
    r.p("El carácter sistemático significa que la calidad se gestiona con un proceso definido, repetible y medible, "
        "no de forma improvisada. El ciclo se repite en cada cambio y en cada sprint:")
    r.tabla("Ciclo sistemático de calidad", ["Paso", "Actividad", "Resultado"], [
        ["1. Planificar", "Definir requisitos, umbrales y estimaciones", "Requisitos y metas"],
        ["2. Construir", "Programar por módulos en sprints", "Código versionado en GitHub"],
        ["3. Verificar", "Pruebas, cobertura y revisión con IA", "Defectos registrados"],
        ["4. Medir", "Calcular métricas de producto, proceso y proyecto", "metrics.json"],
        ["5. Mejorar", "Analizar el informe de calidad y corregir", "Acciones de mejora"]], [3.0, 8.0, 5.5])
    r.p("Como el flujo completo se ejecuta con un solo comando (python run_all.py) y también en GitHub Actions, "
        "cualquier persona obtiene los mismos resultados con el mismo procedimiento.")

    # 6 ------------------------------------------------------------------
    r.h("6. Características de calidad (ISO/IEC 25010)")
    r.p("De las características del modelo de calidad del producto de ISO/IEC 25010 se evalúan dos: idoneidad "
        "funcional y fiabilidad, apoyadas con la mantenibilidad (complejidad ciclomática).")
    r.h("6.1 Idoneidad funcional", 2)
    r.p("Es el grado en que el software proporciona las funciones que satisfacen las necesidades declaradas.")
    r.tabla("Subcaracterísticas de idoneidad funcional", ["Subcaracterística", "Medida", "Resultado"], [
        ["Completitud funcional", "Requisitos cumplidos / requisitos totales", f"{rq['cumplidos']}/{rq['total']} = {rq['porcentaje']} %"],
        ["Corrección funcional", "Pruebas aprobadas / pruebas totales", f"{m['pruebas']['aprobadas']}/{m['pruebas']['total']} = {m['pruebas']['tasa_exito']} %"],
        ["Pertinencia funcional", "Requisitos con al menos una prueba", f"{sum(1 for x in rq['detalle'] if x['pruebas'] > 0)}/{rq['total']}"],
        ["Cobertura de verificación", "Sentencias ejecutadas por las pruebas", f"{co['total']} %"]], [4.5, 7.0, 5.0])
    r.h("6.2 Fiabilidad", 2)
    r.p("Es el grado en que el sistema se mantiene funcionando y se recupera de fallas.")
    r.tabla("Subcaracterísticas de fiabilidad", ["Subcaracterística", "Medida", "Resultado"], [
        ["Madurez", "Densidad de defectos (defectos/KLOC)", f"{d['densidad_kloc']}"],
        ["Tolerancia a fallos", "Prueba: la IA falla y el asistente responde igual", si_no(any(x["id"] == "RNF-01" and x["cumple"] for x in rq["detalle"]))],
        ["Capacidad de recuperación", "Prueba de archivo dañado + MTTR", f"{si_no(any(x['id'] == 'RNF-02' and x['cumple'] for x in rq['detalle']))} / {d['mttr_horas']} h"],
        ["Disponibilidad", "No aplica: es una aplicación de consola, sin servicio permanente", "No medida"]],
        [4.5, 7.0, 5.0])

    # 7 ------------------------------------------------------------------
    r.h("7. Ejecución del programa y salida del código")
    r.p("La salida mostrada es la que produjo calidad/demo.py durante la ejecución de run_all.py; se guardó en "
        "salidas/salida_programa.txt y se convirtió a imagen. El asistente de IA operó en modo: "
        f"{m['modo_ia']}. Los relojes de los tickets de la demostración son simulados para poder medir tiempos de resolución.")
    r.figura(figura_terminal("salida_programa.txt", "term_programa", "python -m calidad.demo"),
             "Salida de la ejecución del programa",
             "Se observan cuatro bloques: el catálogo con precios por m2; tres pedidos con su subtotal, descuento "
             "(0 %, 5 % y 10 %) y total con IVA; las respuestas del asistente a una cotización en texto libre, una "
             "comparación de productos y dos fallas (severidad media y alta); y el cierre de tickets con su tiempo de resolución.")
    try:
        with open(f"{SAL}/demo_datos/pedidos.json", encoding="utf-8") as f:
            pedidos = json.load(f)
        r.tabla("Pedidos registrados en la demostración", ["#", "Cliente", "Tipo", "Subtotal", "Desc.", "Total", "Estado"],
                [[p["id"], p["cliente"], p["cotizacion"]["tipo"], f"${p['cotizacion']['subtotal']:,.2f}",
                  f"{p['descuento_pct']:.0%}", f"${p['total']:,.2f}", p["estado"]] for p in pedidos],
                [1.0, 3.2, 2.5, 3.0, 1.6, 3.2, 2.0], ["c", "l", "l", "r", "c", "r", "c"])
    except FileNotFoundError:
        pass
    tk = m["tickets"]
    r.p(f"Tickets de soporte en la demostración: {tk['total']} creados, {tk['resueltos']} resueltos, "
        f"tiempo promedio de resolución {tk['minutos_promedio']} minutos (reloj simulado).")
    r.figura(figura_terminal("salida_pruebas.txt", "term_pruebas", "python -m pytest --cov=app -v"),
             "Salida de las pruebas automáticas con cobertura",
             f"Cada línea es una prueba con su resultado. Terminaron {m['pruebas']['aprobadas']} de "
             f"{m['pruebas']['total']} aprobadas en {m['pruebas']['tiempo_s']} s. La tabla final muestra el porcentaje "
             f"de sentencias cubiertas por módulo y el total de {co['total']} %.", ancho=14.5)
    r.figura(captura_o_marcador("consola_windows", "Captura de PowerShell ejecutando python run_all.py"),
             "Captura real de la ejecución en Windows",
             "Muestra el comando python run_all.py ejecutándose en PowerShell y su resumen final de métricas; "
             "confirma que el programa corre en el entorno Windows del alumno.")

    # 8 ------------------------------------------------------------------
    r.h("8. Métricas de producto")
    r.p("Miden las características del código y del software ya construido.")
    r.h("8.1 Complejidad ciclomática", 2)
    r.p("Mide el número de rutas independientes del código. Un valor alto indica código difícil de mantener y probar. "
        "Se calcula con radon; se recomienda mantener cada función en 10 o menos.")
    r.formula("CC = decisiones + 1")
    filas = [[k, v["sloc"], v["cc_promedio"], v["cc_maxima"], v["rango_max"], v["indice_mantenibilidad"],
              f"{co['por_archivo'].get(k, {}).get('cobertura', '-')}"] for k, v in pr["modulos"].items()]
    r.tabla("Métricas de producto por módulo", ["Módulo", "SLOC", "CC prom.", "CC máx.", "Rango", "Mantenib.", "Cobertura %"],
            filas, [3.6, 1.8, 2.0, 2.0, 1.8, 2.4, 2.9], ["l", "c", "c", "c", "c", "c", "c"])
    r.figura(figura_complejidad(m), "Funciones con mayor complejidad ciclomática",
             f"Cada barra es una función; la línea roja marca el límite recomendado de 10. La más compleja tiene "
             f"CC = {pr['cc_maxima']} y el promedio del sistema es {pr['cc_promedio']}, por lo que ninguna función supera "
             f"el límite y el código es fácil de probar y mantener." if pr["cc_maxima"] <= 10 else
             f"Cada barra es una función; la línea roja marca el límite de 10. Hay funciones que lo superan y deben dividirse.")
    r.h("8.2 Cobertura de código", 2)
    r.p("Porcentaje de líneas ejecutadas durante las pruebas automatizadas.")
    r.formula("Cobertura = sentencias ejecutadas / sentencias totales x 100")
    r.figura(figura_cobertura(m), "Cobertura de código por módulo",
             f"La cobertura total es {co['total']} %. La línea roja es la meta de 80 %. Los módulos por debajo del 100 % "
             "tienen pocas sentencias sin cubrir, principalmente ramas de manejo de errores poco frecuentes.")
    r.h("8.3 Densidad de defectos", 2)
    r.p("Cantidad de errores en relación con el tamaño del software, medido en miles de líneas de código (KLOC).")
    r.formula(f"Densidad = {d['total']} defectos / {pr['kloc']} KLOC = {d['densidad_kloc']} defectos/KLOC")
    r.tabla("Densidad de defectos por módulo", ["Módulo", "Defectos", "SLOC", "Densidad (def/KLOC)"],
            [[k, v["defectos"], pr["modulos"][k]["sloc"], v["densidad"]] for k, v in d["por_modulo"].items()],
            [4.5, 3.0, 3.0, 4.5], ["l", "c", "c", "c"])
    r.figura(figura_defectos(m), "Defectos por fase y densidad por módulo",
             f"A la izquierda se ve en qué fase se detectó cada defecto: {d['por_fase']['revision']} en revisión, "
             f"{d['por_fase']['pruebas']} en pruebas y {d['por_fase']['produccion']} en producción. A la derecha, la densidad por "
             "módulo permite ver dónde se concentran los defectos y priorizar la mejora.")
    r.h("8.4 Interpretación asistida por IA", 2)
    ll = m.get("ia_llamadas", {"ok": 0, "fallidas": 0})
    r.p(f"Modo de IA: {m['modo_ia']}. Llamadas a la IA durante el cálculo de métricas: {ll['ok']} exitosas, "
        f"{ll['fallidas']} fallidas.", cursiva=True)
    r.p(m["interpretacion_ia"])

    # 9 ------------------------------------------------------------------
    r.h("9. Métricas del proceso")
    r.p("Miden la eficiencia de los flujos de trabajo, las revisiones y las fases de prueba.")
    r.tabla("Resumen de métricas de proceso", ["Métrica", "Fórmula", "Resultado"], [
        ["Tiempo medio de detección (MTTD)", "Promedio (detección - introducción)", f"{d['mttd_horas']} h"],
        ["Tiempo medio de reparación (MTTR)", "Promedio (resolución - detección)", f"{d['mttr_horas']} h"],
        ["Eficacia de las pruebas", "Críticos en pruebas / (críticos en pruebas + en producción)",
         f"{d['criticos_pruebas']}/{d['criticos_pruebas'] + d['criticos_produccion']} = {d['eficacia_pruebas']} %"]],
        [5.3, 7.2, 4.0])
    r.tabla("Bitácora de defectos", ["ID", "Módulo", "Severidad", "Fase", "Origen", "MTTD h", "MTTR h"],
            [[b["id"], b["modulo"], b["severidad"], b["fase_deteccion"], b["origen"], b["mttd"], b["mttr"]] for b in d["bitacora"]],
            [1.4, 3.2, 2.5, 2.7, 2.2, 2.2, 2.3], ["c", "l", "c", "c", "c", "c", "c"])
    r.figura(figura_tiempos(m), "Tiempo de detección y reparación por defecto",
             f"Las barras azules son las horas desde que se introdujo el defecto hasta que se detectó; las naranjas, "
             f"hasta que se reparó. El promedio de detección es {d['mttd_horas']} h y el de reparación {d['mttr_horas']} h. "
             "Los defectos que tardaron más en detectarse son los que se encontraron en pruebas o producción, lo que justifica la revisión temprana.")
    r.p(f"Eficacia de las pruebas: de los defectos críticos, {d['criticos_pruebas']} se detectaron en pruebas y "
        f"{d['criticos_produccion']} llegaron a producción, por lo tanto la eficacia es {d['eficacia_pruebas']} %.")
    r.p(f"Soporte posventa: en la demostración se resolvieron {tk['resueltos']} de {tk['total']} tickets en un promedio de "
        f"{tk['minutos_promedio']} minutos (reloj simulado).")

    # 10 -----------------------------------------------------------------
    r.h("10. Métricas del proyecto")
    r.p("Analizan la gestión del equipo y los plazos con la metodología ágil (sprints).")
    r.h("10.1 Eficacia de la revisión", 2)
    r.formula(f"Eficacia de revisión = {d['por_fase']['revision']} / {d['total']} x 100 = {d['eficacia_revision']} %")
    r.p(f"Es el porcentaje de defectos encontrados en las revisiones de código, antes de las pruebas. "
        f"De ellos, {d['por_origen']['IA']} de {d['total']} fueron señalados por la IA y {d['por_origen']['humano']} por el desarrollador.")
    r.h("10.2 Desviación de tiempo y esfuerzo", 2)
    r.formula("Desviación % = (horas reales - horas estimadas) / horas estimadas x 100")
    r.tabla("Desviación por módulo y densidad de defectos", ["Módulo", "Sprint", "Estimadas h", "Reales h", "Desviación %", "Def/KLOC"],
            [[x["modulo"], x["sprint"], x["estimadas"], x["reales"], x["desviacion_pct"], x["densidad"] if x["densidad"] is not None else "n/a"]
             for x in py["modulos"]] + [["TOTAL", "-", py["horas_estimadas"], py["horas_reales"], py["desviacion_pct"], d["densidad_kloc"]]],
            [4.2, 1.6, 2.7, 2.4, 2.9, 2.7], ["l", "c", "c", "c", "c", "c"])
    r.figura(figura_desviacion(m), "Horas estimadas contra horas reales por módulo",
             f"Las barras grises son lo planeado al inicio del sprint y las azules lo realmente invertido. El proyecto total "
             f"se desvió {py['desviacion_pct']} %. El módulo soporte_ia fue el de mayor esfuerzo por la integración con IA y su tolerancia a fallos.")
    r.tabla("Resumen por sprint", ["Sprint", "Estimadas h", "Reales h", "Desviación %", "Defectos"],
            [[k, v["estimadas"], v["reales"], v["desviacion_pct"], v["defectos"]] for k, v in py["sprints"].items()],
            [2.5, 3.5, 3.5, 3.5, 3.5], ["c"] * 5)

    # 11 -----------------------------------------------------------------
    r.h("11. Técnicas de estimación")
    je = es["juicio_expertos"]
    r.h("11.1 Juicio de expertos", 2)
    r.p("Varias personas con experiencia estiman cada módulo y se toma el promedio como consenso. Expertos: " + ", ".join(je["expertos"]) + ".")
    r.tabla("Estimación por juicio de expertos (horas)", ["Módulo"] + je["expertos"] + ["Promedio"],
            [[k] + v["valores"] + [v["promedio"]] for k, v in je["por_modulo"].items()] + [["TOTAL", "", "", "", je["total"]]],
            [4.3, 3.2, 3.5, 2.8, 2.7], ["l", "c", "c", "c", "c"])
    an = es["analoga"]
    r.h("11.2 Estimación análoga", 2)
    r.p(f"Se compara con un proyecto anterior parecido: {an['proyecto_anterior']}.")
    r.tabla("Estimación análoga", ["Dato", "Valor"], [
        ["Horas del proyecto anterior", an["horas_anterior"]], ["SLOC del proyecto anterior", an["sloc_anterior"]],
        ["SLOC del proyecto actual (medido)", an["sloc_actual"]], ["Horas por SLOC", an["horas_por_sloc"]],
        ["Factor de ajuste", an["factor_ajuste"]], ["Estimación (horas)", an["horas"]]], [8.0, 5.0], ["l", "c"])
    r.formula("Estimación = horas anteriores / SLOC anterior x SLOC actual x factor de ajuste")
    r.p(an["justificacion"])
    tp = es["tres_puntos"]
    r.h("11.3 Estimación de tres puntos (PERT)", 2)
    r.formula("E = (O + 4M + P) / 6        sigma = (P - O) / 6")
    r.tabla("Estimación de tres puntos por módulo (horas)", ["Módulo", "O", "M", "P", "E", "sigma"],
            [[k, v["O"], v["M"], v["P"], v["esperado"], v["sigma"]] for k, v in tp["por_modulo"].items()] +
            [["TOTAL", "", "", "", tp["total"], tp["sigma_total"]]], [4.6, 2.0, 2.0, 2.0, 3.0, 3.0], ["l"] + ["c"] * 5)
    r.p(f"Esfuerzo esperado total: {tp['total']} h. Con 95 % de confianza (+/- 2 sigma) el rango es de "
        f"{tp['rango_95'][0]} a {tp['rango_95'][1]} h; las horas reales ({py['horas_reales']:.0f} h) "
        f"{'caen dentro' if tp['rango_95'][0] <= py['horas_reales'] <= tp['rango_95'][1] else 'quedan fuera'} de ese rango.")
    pf = es["puntos_funcion"]
    r.h("11.4 Puntos de función", 2)
    r.p("Se cuenta la funcionalidad desde la vista del usuario: entradas (EI), salidas (EO), consultas (EQ), archivos "
        "internos (ILF) e interfaces externas (EIF), con pesos IFPUG según su complejidad.")
    r.tabla("Componentes funcionales", ["Componente", "Tipo", "Complejidad", "Peso"],
            [[c["nombre"], c["tipo"], c["complejidad"], c["peso"]] for c in pf["detalle"]], [8.0, 2.0, 3.5, 2.0], ["l", "c", "c", "c"])
    r.tabla("Resumen de puntos de función", ["Tipo", "Cantidad", "Puntos"],
            [[k, v["cantidad"], v["puntos"]] for k, v in pf["resumen"].items()] + [["UFP (sin ajustar)", "", pf["ufp"]]], [6.0, 4.0, 4.0], ["l", "c", "c"])
    r.formula(f"VAF = 0.65 + 0.01 x {pf['suma_gsc']} = {pf['vaf']}     PF = {pf['ufp']} x {pf['vaf']} = {pf['pf']}")
    r.p(f"Con una productividad de {pf['horas_por_pf']} h por punto de función, el esfuerzo estimado es {pf['horas']} h.")
    r.h("11.5 Comparativo de técnicas", 2)
    r.tabla("Estimaciones contra esfuerzo real", ["Técnica", "Estimación h", "Real h", "Diferencia %"],
            [[c["tecnica"], c["horas"], c["real"], c["desviacion_pct"]] for c in es["comparativo"]], [6.0, 3.5, 3.0, 3.5], ["l", "c", "c", "c"])
    mejor = min(es["comparativo"], key=lambda c: abs(c["desviacion_pct"]))
    r.figura(figura_estimaciones(m), "Comparación de las cuatro técnicas contra las horas reales",
             f"La línea roja son las horas reales. La técnica más cercana fue {mejor['tecnica']} ({mejor['desviacion_pct']} %). "
             "Las diferencias muestran que cada técnica captura riesgos distintos y conviene combinarlas.")

    # 12 -----------------------------------------------------------------
    r.h("12. Revisión de código asistida por IA")
    r.p(f"Modo de revisión: {ri['modo']}. Se revisaron los módulos de app/ y se obtuvieron {ri['total']} hallazgos "
        "de mantenibilidad (distintos de la bitácora de defectos).")
    r.tabla("Hallazgos de la revisión", ["Archivo", "Línea", "Regla", "Severidad", "Descripción"],
            [[h["archivo"], h["linea"], h["regla"], h["severidad"], h["descripcion"]] for h in ri["hallazgos"][:12]] or
            [["-", "-", "-", "-", "Sin hallazgos"]], [3.0, 1.5, 1.6, 2.2, 8.2], ["l", "c", "c", "c", "l"])
    r.p("Cada hallazgo incluye una sugerencia de prueba (en salidas/metrics.json) que alimenta nuevos casos de prueba y "
        "refuerza el enfoque preventivo.")

    # 13 -----------------------------------------------------------------
    r.h("13. Informe de calidad")
    r.tabla("Ficha del informe", ["Campo", "Valor"], [
        ["Proyecto", "Sistema de ventas de letreros luminosos con soporte IA"], ["Fecha de ejecución", m["generado"].replace("T", " ")],
        ["Entorno", f"{m['sistema']} / Python {m['python']}"], ["Modo de IA", m["modo_ia"]],
        ["Tamaño", f"{pr['sloc']} SLOC, {pr['funciones_total']} funciones"]], [4.5, 12.0])
    r.tabla("Cumplimiento de criterios de calidad", ["Criterio", "Característica", "Meta", "Resultado", "Estado"],
            [[c["criterio"], c["caracteristica"], c["meta"], c["resultado"], "Cumple" if c["cumple"] else "No cumple"] for c in ca["criterios"]],
            [5.2, 3.3, 2.5, 2.5, 3.0], ["l", "l", "c", "c", "c"])
    r.p(f"Grado de cumplimiento de calidad: {ca['cumplen']} de {ca['total']} criterios = {ca['grado_cumplimiento']} %.", True)
    r.tabla("Riesgos y acciones de mejora", ["Riesgo observado", "Acción de mejora"], [
        [f"{d['criticos_produccion']} defecto(s) crítico(s) llegó a producción", "Agregar pruebas de severidad en el asistente antes de liberar"],
        ["Datos en archivos JSON", "Migrar a SQLite en la Fase 2 con el mismo repositorio"],
        ["Dependencia de un proveedor externo de IA", "Mantener el modo local y la prueba de tolerancia a fallos"],
        ["La disponibilidad no se mide", "Agregar monitoreo cuando exista un servicio web"]], [7.0, 9.5])

    # 14 -----------------------------------------------------------------
    r.h("14. GitHub, repositorio y CI/CD")
    r.p(f"Repositorio: {cfg['repositorio']}")
    r.tabla("Estructura del repositorio", ["Ruta", "Propósito"], [
        ["app/", "Módulos del sistema (catálogo, cotizador, ventas, soporte IA, repositorio)"],
        ["tests/", "Pruebas automáticas con pytest"], ["calidad/", "Métricas, estimaciones, revisión IA y datos del proyecto"],
        ["run_all.py", "Flujo completo con un solo comando"], ["generar_reporte.py", "Genera este reporte desde metrics.json"],
        [".github/workflows/ci.yml", "Integración continua"], ["Dockerfile", "Ejecución en contenedor"]], [5.5, 11.0])
    r.p("El workflow ci.yml instala dependencias, ejecuta pytest con un mínimo de 80 % de cobertura (el pipeline falla si no se cumple), "
        "corre run_all.py y publica metrics.json y el reporte como artefactos.")
    r.figura(captura_o_marcador("github_repositorio", "Captura de la página principal del repositorio en GitHub"),
             "Repositorio en GitHub",
             "Muestra los archivos y carpetas del proyecto y el historial de commits, evidencia del control de versiones.")
    r.figura(captura_o_marcador("github_actions", "Captura de la pestaña Actions con la ejecución en verde"),
             "Ejecución del workflow en GitHub Actions",
             "Muestra que el pipeline se ejecutó tras un push y que sus pasos (instalación, pruebas, métricas) terminaron correctamente.")

    # 15 -----------------------------------------------------------------
    r.h("15. Conclusiones")
    r.p(f"El sistema cumple {rq['cumplidos']} de {rq['total']} requisitos y {ca['cumplen']} de {ca['total']} criterios de calidad "
        f"({ca['grado_cumplimiento']} %), con {co['total']} % de cobertura y complejidad máxima de {pr['cc_maxima']}.")
    r.p(f"El flujo DevOps permite repetir la validación sin esfuerzo manual: un comando ejecuta el programa, las pruebas, las "
        f"métricas y el reporte. Esto garantiza que lo documentado coincide con lo ejecutado y reduce el riesgo de errores humanos.")
    r.p(f"La IA aportó en tres puntos: atiende al cliente y diagnostica fallas en el soporte, revisa el código y interpreta las métricas. "
        f"Por diseño, los precios y diagnósticos los calcula código verificable y la IA solo redacta, y el sistema sigue funcionando si la IA no está disponible.")
    r.p(f"En el proceso, el MTTD de {d['mttd_horas']} h frente al MTTR de {d['mttr_horas']} h muestra que el reto principal es detectar antes, no reparar; "
        f"por eso se refuerza la revisión temprana ({d['eficacia_revision']} % de eficacia). La desviación de esfuerzo fue de {py['desviacion_pct']} %, "
        "y las cuatro técnicas de estimación ofrecieron un rango útil para planear mejor el siguiente sprint.")
    r.p("Como trabajo futuro (Fase 2), se reemplazará el repositorio JSON por SQLite, conservando la lógica, las pruebas y el pipeline, "
        "y se recalcularán las métricas con el mismo comando.")

    # Anexo --------------------------------------------------------------
    r.h("Anexo A. Cómo ejecutar en Windows (PowerShell)")
    r.tabla("Comandos", ["Paso", "Comando"], [
        ["Crear entorno", "python -m venv .venv ; .\\.venv\\Scripts\\Activate.ps1"],
        ["Instalar", "pip install -r requirements.txt"], ["Ejecutar todo", "python run_all.py"],
        ["Usar el programa", "python main.py"], ["Solo pruebas", "python -m pytest --cov=app -v"],
        ["IA con Gemini (opcional)", "$env:GEMINI_API_KEY = 'tu-clave' ; python run_all.py"],
        ["IA con Claude (opcional)", "$env:ANTHROPIC_API_KEY = 'tu-clave' ; python run_all.py"]], [4.5, 12.0])
    ruta = "Reporte_Letreros_Luminosos.docx"
    r.doc.save(ruta)
    return ruta


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    with open(f"{SAL}/metrics.json", encoding="utf-8") as f:
        print(generar(json.load(f)))
