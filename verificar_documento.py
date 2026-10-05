"""Comprueba que las cifras clave de metrics.json aparezcan en el reporte Word.

Uso:  python verificar_documento.py
Primero corre `python run_all.py` para generar el reporte y las métricas.
"""
import json
import re
import sys

from docx import Document

REPORTE = "Reporte_Letreros_Luminosos.docx"


def texto_del_documento(ruta):
    doc = Document(ruta)
    partes = [p.text for p in doc.paragraphs]
    for tabla in doc.tables:
        for fila in tabla.rows:
            partes += [celda.text for celda in fila.cells]
    return "\n".join(partes)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    with open("salidas/metrics.json", encoding="utf-8") as f:
        m = json.load(f)
    texto = texto_del_documento(REPORTE)
    pr, d, py, co = m["producto"], m["defectos"], m["proyecto"], m["cobertura"]
    es = m["estimaciones"]
    revisiones = [
        ("Pruebas totales", m["pruebas"]["total"]), ("Pruebas aprobadas", m["pruebas"]["aprobadas"]),
        ("Cobertura total %", co["total"]), ("Líneas de código (SLOC)", pr["sloc"]),
        ("Complejidad promedio", pr["cc_promedio"]), ("Complejidad máxima", pr["cc_maxima"]),
        ("Densidad de defectos", d["densidad_kloc"]), ("MTTD (h)", d["mttd_horas"]),
        ("MTTR (h)", d["mttr_horas"]), ("Eficacia de pruebas %", d["eficacia_pruebas"]),
        ("Eficacia de revisión %", d["eficacia_revision"]),
        ("Desviación total de esfuerzo %", py["desviacion_pct"]),
        ("Requisitos cumplidos %", m["requisitos"]["porcentaje"]),
        ("Grado de cumplimiento %", m["calidad"]["grado_cumplimiento"]),
        ("Juicio de expertos (h)", es["juicio_expertos"]["total"]),
        ("Estimación análoga (h)", es["analoga"]["horas"]),
        ("Tres puntos (h)", es["tres_puntos"]["total"]),
        ("Puntos de función (PF)", es["puntos_funcion"]["pf"]),
        ("Horas por puntos de función", es["puntos_funcion"]["horas"]),
        ("Modo de IA", m["modo_ia"]),
        ("Pruebas de IA aprobadas", f"{m['pruebas_ia']['pasan']} de {m['pruebas_ia']['total']}"),
    ]
    fallos = 0
    print(f"Verificando {REPORTE} contra salidas/metrics.json\n")
    for nombre, valor in revisiones:
        patron = r"(?<![\d.])" + re.escape(str(valor)) + r"(?![\d])"
        ok = re.search(patron, texto) is not None
        fallos += not ok
        print(f"[{'OK' if ok else 'FALTA'}] {nombre}: {valor}")
    marcadores = sorted(set(re.findall(r"\[[^\]\n]{3,60}\]", texto)))
    marcadores = [x for x in marcadores if "[Pega" not in x]
    if marcadores:
        print("\nAviso: quedan textos entre corchetes sin llenar:", marcadores)
    print(f"\nResultado: {len(revisiones) - fallos} de {len(revisiones)} cifras coinciden.")
    print("Recuerda: esto confirma que las cifras están en el documento; "
          "revisa también que los textos tengan sentido con tus datos.")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
